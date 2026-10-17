"""Erzeugt die Excel-Dateien für das Netzlaufwerk.

dateien/Messprotokoll.xlsx  -> wird nach dem VBA-Import als Messprotokoll.xlsm gespeichert
dateien/Stammdaten.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import iso286_daten as D

BASIS = Path(__file__).resolve().parent.parent
ZIEL = BASIS / "dateien"

PP_ERSTE = 8
PP_MAX = 1000
LETZTE = PP_ERSTE + PP_MAX - 1

BLAU = "1F4E79"
HELLBLAU = "DDEBF7"
GRAU = "F2F2F2"
GELB_EINGABE = "FFF2CC"
fett = Font(bold=True)
weiss_fett = Font(bold=True, color="FFFFFF")
kopf_fill = PatternFill("solid", fgColor=BLAU)
eingabe_fill = PatternFill("solid", fgColor=GELB_EINGABE)
grau_fill = PatternFill("solid", fgColor=GRAU)
duenn = Side(style="thin", color="BFBFBF")
rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)

ARTEN = ["Maß", "Passung", "Form/Lage", "Oberfläche", "Attribut"]


def titel(ws, text, unter=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=BLAU)
    if unter:
        ws["A2"] = unter
        ws["A2"].font = Font(italic=True, color="595959")


def blatt_start(wb):
    ws = wb.active
    ws.title = "Start"
    titel(ws, "Messprotokoll · Metalltechnik", "Messdatenverwaltung nach ISO 9001")
    for z, text in enumerate(["Version", "Basisordner", "Angemeldet als", "Rollen"], start=4):
        ws.cell(z, 2, text).font = fett
    ws["B9"] = "Prüfpläne"
    ws["B9"].font = Font(bold=True, size=12, color=BLAU)
    ws["B19"] = "Verwaltung"
    ws["B19"].font = Font(bold=True, size=12, color=BLAU)
    ws["E9"] = "So geht's"
    ws["E9"].font = Font(bold=True, size=12, color=BLAU)
    hilfe = [
        "Prüfplan laden: Artikelnummer eingeben, Stand wählen.",
        "Neuer Prüfplan: für Artikel, die noch keinen Prüfplan haben.",
        "Neuer Zeichnungsstand: kopiert den geladenen Stand.",
        "   Nur Änderungen eintragen – geänderte Merkmale werden",
        "   beim Speichern gelb, neue grün markiert.",
        "   Der alte Stand wird abgeschlossen (nur lesen).",
        "Prüfplan speichern: prüft alle Zeilen und legt die Datei ab.",
        "",
        "Merkmalsarten im Prüfplan:",
        "   Maß – Nennmaß + Abmaße oder Allgemeintoleranz (2768-m)",
        "   Passung – Nennmaß + ISO-Passung (H7, g6 …)",
        "   Form/Lage – Toleranz im Feld 'Abmaß oben'",
        "   Oberfläche – Höchstwert (z. B. Ra 1,6) in 'Abmaß oben'",
        "   Attribut – nur i.O./n.i.O. (z. B. Gewindelehre)",
    ]
    for i, t in enumerate(hilfe):
        ws.cell(10 + i, 5, t)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 3
    ws.column_dimensions["E"].width = 62
    for z in range(10, 24):
        ws.row_dimensions[z].height = 18
    ws.sheet_view.showGridLines = False


def blatt_pruefplan(wb):
    ws = wb.create_sheet("Prüfplan")
    ws["A1"] = "Prüfplan"
    ws["A1"].font = Font(bold=True, size=16, color=BLAU)
    links = [("B2", "Artikelnummer"), ("B3", "Bezeichnung"), ("B4", "Zeichnungsstand"), ("B5", "Basis-Stand")]
    rechts = [("F2", "Erstellt von"), ("F3", "Erstellt am"), ("F4", "Status"), ("F5", "Entfallene Merkmale")]
    for adr, text in links + rechts:
        ws[adr] = text
        ws[adr].font = fett
        ws[adr].alignment = Alignment(horizontal="right")
    for adr in ["C2", "C3", "C4"]:
        ws[adr].fill = eingabe_fill
        ws[adr].border = rahmen
    for adr in ["C5", "G2", "G3", "G4", "G5"]:
        ws[adr].fill = grau_fill
        ws[adr].border = rahmen
    ws["G3"].number_format = "DD.MM.YYYY HH:MM"
    ws["A6"] = ("Gelb = Eingabe · Grau = wird berechnet · Spalte N zeigt Eingabefehler. "
                "Passung: Art 'Passung' + z. B. H7 · Allgemeintoleranz: Art 'Maß' + z. B. 2768-m")
    ws["A6"].font = Font(italic=True, color="595959", size=9)

    spalten = [
        ("Nr", 7), ("Merkmal", 28), ("Art", 12), ("Nennmaß", 10), ("Passung / Toleranz", 14),
        ("Abmaß oben", 11), ("Abmaß unten", 11), ("OGW", 11), ("UGW", 11),
        ("Messmittel", 18), ("Kritisch", 9), ("Änderung", 10), ("Bemerkung", 28), ("Prüfung", 40),
    ]
    for i, (text, breite) in enumerate(spalten, start=1):
        c = ws.cell(7, i, text)
        c.font = weiss_fett
        c.fill = kopf_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = breite
    ws.row_dimensions[7].height = 30

    for z in range(PP_ERSTE, LETZTE + 1):
        for sp in (4, 6, 7):
            ws.cell(z, sp).number_format = "0.000"
        for sp in (8, 9):
            ws.cell(z, sp).number_format = "0.000"
            ws.cell(z, sp).fill = grau_fill
        ws.cell(z, 12).fill = grau_fill
        ws.cell(z, 14).font = Font(color="C00000")

    dv_art = DataValidation(type="list", formula1='"' + ",".join(ARTEN) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Art",
                            error="Bitte eine Art aus der Liste wählen.")
    dv_art.add(f"C{PP_ERSTE}:C{LETZTE}")
    dv_krit = DataValidation(type="list", formula1='"ja,nein"', allow_blank=True)
    dv_krit.add(f"K{PP_ERSTE}:K{LETZTE}")
    ws.add_data_validation(dv_art)
    ws.add_data_validation(dv_krit)

    ws.conditional_formatting.add(
        f"A{PP_ERSTE}:M{LETZTE}",
        FormulaRule(formula=[f'$N{PP_ERSTE}<>""'], font=Font(color="C00000")))
    ws.freeze_panes = f"B{PP_ERSTE}"
    ws.auto_filter.ref = f"A7:N{LETZTE}"


def blatt_rechner(wb):
    ws = wb.create_sheet("ISO-Rechner")
    titel(ws, "ISO-Passungsrechner", "Nennmaß und Passung eingeben – Abmaße nach ISO 286")
    werte = [("B4", "Nennmaß (mm)", "C4", 20), ("B5", "Passung", "C5", "H7")]
    for adr, text, zadr, wert in werte:
        ws[adr] = text
        ws[adr].font = fett
        ws[zadr] = wert
        ws[zadr].fill = eingabe_fill
        ws[zadr].border = rahmen
    for adr, text in [("B7", "Abmaß oben (mm)"), ("B8", "Abmaß unten (mm)"),
                      ("B9", "Höchstmaß (mm)"), ("B10", "Mindestmaß (mm)")]:
        ws[adr] = text
        ws[adr].font = fett
        z = "C" + adr[1:]
        ws[z].fill = grau_fill
        ws[z].border = rahmen
        ws[z].number_format = "0.000"
    ws["B12"] = ("Unterstützt: Bohrungen D E F G H JS K M N P R S, Wellen d e f g h js k m n p r s, "
                 "IT1–IT16, Nennmaß bis 500 mm.")
    ws["B12"].font = Font(italic=True, color="595959")
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 18
    ws.sheet_view.showGridLines = False


def blatt_iso_tabelle(wb):
    ws = wb.create_sheet("ISO286")
    titel(ws, "ISO 286-1 Grundtabellen (µm)", "Nur zum Nachschlagen – gerechnet wird mit denselben Daten im Programm.")
    bereiche = [f"{a}–{b}" for a, b in zip([0] + D.HAUPT_BIS[:-1], D.HAUPT_BIS)]
    zeile = 4

    def tabelle(ueberschrift, kopf, reihen):
        nonlocal zeile
        ws.cell(zeile, 1, ueberschrift).font = Font(bold=True, color=BLAU)
        zeile += 1
        for i, k in enumerate(["Nennmaß über–bis mm"] + kopf, start=1):
            c = ws.cell(zeile, i, k)
            c.font = weiss_fett
            c.fill = kopf_fill
        zeile += 1
        for name, werte in reihen:
            ws.cell(zeile, 1, name).font = fett
            for i, w in enumerate(werte, start=2):
                ws.cell(zeile, i, w)
            zeile += 1
        zeile += 1

    tabelle("Grundtoleranzen IT", bereiche, [(f"IT{g}", w) for g, w in D.IT.items()])
    tabelle("Wellen: oberes Abmaß es", bereiche, list(D.WELLE_ES.items()))
    tabelle("Wellen: unteres Abmaß ei (k nur IT4–IT7)", bereiche, list(D.WELLE_EI.items()))
    fein = [f"{a}–{b}" for a, b in zip([0] + D.FEIN_BIS[:-1], D.FEIN_BIS)]
    tabelle("Wellen: unteres Abmaß ei (Feinbereiche)", fein, list(D.WELLE_EI_FEIN.items()))
    ws.cell(zeile, 1, "Bohrungen: D–H spiegelbildlich zu d–h (EI = −es); "
                      "K, M, N bis IT8 und P, R, S bis IT7: ES = −ei + Δ mit Δ = IT(n) − IT(n−1).").font = Font(italic=True)
    ws.column_dimensions["A"].width = 22


def blatt_system(wb):
    ws = wb.create_sheet("System")
    for z, text in enumerate(["Geladene Prüfplan-Datei", "Zeitstempel der Datei",
                              "Basisordner (leer = automatisch)", "Version"], start=1):
        ws.cell(z, 1, text).font = fett
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80
    ws.sheet_state = "hidden"


def messprotokoll():
    wb = Workbook()
    blatt_start(wb)
    blatt_pruefplan(wb)
    blatt_rechner(wb)
    blatt_iso_tabelle(wb)
    blatt_system(wb)
    wb.save(ZIEL / "Messprotokoll.xlsx")


def stammdaten():
    wb = Workbook()
    ws = wb.active
    ws.title = "Personen"
    kopf = [("Benutzername", 18), ("Name", 26), ("Prüfer", 10), ("Freigeber", 11), ("Stammdaten", 12)]
    for i, (t, b) in enumerate(kopf, start=1):
        c = ws.cell(1, i, t)
        c.font = weiss_fett
        c.fill = kopf_fill
        ws.column_dimensions[c.column_letter].width = b
    for z in range(2, 302):
        for sp in range(1, 6):
            ws.cell(z, sp).border = rahmen
    dv = DataValidation(type="list", formula1='"x"', allow_blank=True)
    dv.add("C2:E301")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A2"

    mm = wb.create_sheet("Messmittel")
    kopf = [("Nr", 10), ("Bezeichnung", 26), ("Art", 18), ("Messbereich", 14), ("Auflösung", 11),
            ("Standort", 16), ("Letzte Kalibrierung", 14), ("Intervall (Monate)", 12),
            ("Nächste Kalibrierung", 14), ("Status", 12)]
    for i, (t, b) in enumerate(kopf, start=1):
        c = mm.cell(1, i, t)
        c.font = weiss_fett
        c.fill = kopf_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        mm.column_dimensions[c.column_letter].width = b
    for z in range(2, 1002):
        mm.cell(z, 7).number_format = "DD.MM.YYYY"
        mm.cell(z, 9, f'=IF(OR(G{z}="",H{z}=""),"",EDATE(G{z},H{z}))').number_format = "DD.MM.YYYY"
        mm.cell(z, 10, f'=IF(I{z}="","",IF(I{z}<TODAY(),"abgelaufen",IF(I{z}-TODAY()<=30,"bald fällig","gültig")))')
    mm.freeze_panes = "A2"
    mm.conditional_formatting.add("J2:J1001", FormulaRule(formula=['$J2="abgelaufen"'], font=Font(color="C00000", bold=True)))
    mm.conditional_formatting.add("J2:J1001", FormulaRule(formula=['$J2="bald fällig"'], font=Font(color="C65911", bold=True)))

    an = wb.create_sheet("Anleitung")
    titel(an, "Stammdaten · Anleitung")
    zeilen = [
        "Personen",
        "  Benutzername = Windows-Anmeldename (wie beim PC-Login, ohne Domäne).",
        "  In den Spalten Prüfer / Freigeber / Stammdaten ein x setzen.",
        "  Solange niemand die Rolle Stammdaten hat, läuft das Programm im Einrichtungsmodus:",
        "  dann darf jeder Prüfpläne anlegen. Bitte zuerst die Stammdaten-Gruppe eintragen!",
        "",
        "Messmittel (wird in Phase 4 an die Erfassung angebunden)",
        "  Letzte Kalibrierung + Intervall ergeben Nächste Kalibrierung und Status.",
        "",
        "Diese Datei darf nur die Stammdaten-Gruppe ändern (Ordnerrechte vom IT setzen lassen).",
    ]
    for i, t in enumerate(zeilen, start=3):
        an.cell(i, 1, t).font = fett if t and not t.startswith(" ") else Font()
    an.column_dimensions["A"].width = 100
    wb.save(ZIEL / "Stammdaten.xlsx")


if __name__ == "__main__":
    ZIEL.mkdir(exist_ok=True)
    messprotokoll()
    stammdaten()
    print("Excel-Dateien erzeugt:", ", ".join(p.name for p in sorted(ZIEL.glob("*.xlsx"))))
