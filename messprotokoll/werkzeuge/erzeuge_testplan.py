"""Erzeugt TESTPLAN.md und dateien/Testliste.xlsx (Abhakliste für den großen Test).

Die erwarteten Cpk-Werte werden aus den Testdaten berechnet.
"""
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import protokoll_csv as P

BASIS = Path(__file__).resolve().parent.parent
TD = BASIS / "dateien" / "testdaten"


def cpk_tabelle():
    werte = {}
    info = {}
    for f in sorted((TD / "Daten/Archiv/2026/DEMO-100").glob("*.csv")):
        for _k, m in P.lesen(f):
            if m["Art"] == "Attribut" or not m["Istwert"]:
                continue
            werte.setdefault(m["Nr"], []).append(float(m["Istwert"]))
            info[m["Nr"]] = (m["Merkmal"], float(m["OGW"]) if m["OGW"] else None, float(m["UGW"]) if m["UGW"] else None)
    zeilen = []
    for nr in sorted(werte, key=int):
        w = werte[nr]
        n = len(w)
        mw = sum(w) / n
        s = math.sqrt(sum((x - mw) ** 2 for x in w) / (n - 1))
        name, ogw, ugw = info[nr]
        seiten = [v for v in ((ogw - mw) / (3 * s) if ogw is not None else None,
                              (mw - ugw) / (3 * s) if ugw is not None else None) if v is not None]
        cpk = min(seiten)
        txt = "sehr gut" if cpk >= 1.67 else "fähig" if cpk >= 1.33 else "bedingt fähig" if cpk >= 1 else "nicht fähig"
        zeilen.append((nr, name, n, cpk, txt))
    return zeilen


def schritte(cpk):
    c = {nr: (name, v, t) for nr, name, _n, v, t in cpk}
    k = lambda nr: f"{c[nr][1]:.2f}".replace(".", ",")  # noqa: E731
    return [
        ("A Installation", "ZIP herunterladen, Rechtsklick → Eigenschaften → „Zulassen“ → OK, dann entpacken.", "Ordner Messprotokoll_1.0.0 mit INSTALLIEREN.bat"),
        ("A Installation", "INSTALLIEREN.bat doppelklicken. Basisordner z. B. C:\\Messdaten_Test, alle Fragen mit J.", "„Installation abgeschlossen“, Programm öffnet sich"),
        ("A Installation", "Falls gelbe Leiste „Makros deaktiviert“: Inhalt aktivieren.", "Keine Leiste mehr bzw. Makros aktiv"),
        ("B Start", "Startblatt ansehen.", "Version 1.0.0, Basisordner, dein Name, Rollen „Prüfer, Freigeber, Stammdaten“, Hinweis „Messmittel: 2 gesperrt, 1 bald fällig“"),
        ("B Start", "„Selbsttest“ klicken.", "„Selbsttest bestanden: alle 98 Prüfungen in Ordnung“"),
        ("B Start", "„Anmeldung prüfen“ klicken.", "Name, Rollen und Basisordner werden angezeigt"),
        ("C Messmittel", "„Messmittel“ klicken.", "15 Messmittel; MS-99 rot „abgelaufen“, RL-02 rot „keine Kalibrierdaten“, MS-03 gelb „bald fällig“, Rest grün"),
        ("D Auswertung", "Blatt Auswertung: C3 = DEMO-100, „Auswerten“.", "Protokolle 30 · i.O. 27 · Sonderfreigabe 2 · Ausschuss 1; 14 Merkmale"),
        ("D Auswertung", "Cpk-Spalte prüfen.", f"Nr 1 ≈ {k('1')} grün · Nr 2 ≈ {k('2')} rot · Nr 3 ≈ {k('3')} rot · Nr 10 ≈ {k('10')} grün · Nr 13 ≈ {k('13')} gelb"),
        ("D Auswertung", "C7 = 2, „Verlauf anzeigen“.", "Liniendiagramm mit 30 Punkten, rote Linien OGW 20,015 / UGW 20,002, graue Mittelwertlinie"),
        ("E Prüfplan", "„Prüfplan laden“ → DEMO-100 → Stand A.", "15 Merkmale mit OGW/UGW; Hinweis „Stand A ist abgeschlossen (nur lesen)“"),
        ("E Prüfplan", "„Neuer Stand“ → B. Merkmal 1: Abmaß oben 0,15. Zeile 15 komplett leeren. In die nächste freie Zeile: 16 | Ø Zapfen | Passung | 12 | h7. „Speichern“, Frage „aktueller Stand?“ mit Ja.", "Nr 1 gelb „geändert“, Nr 16 grün „neu“, Entfallene Merkmale: 15; OGW Nr 16 = 12,000, UGW = 11,982"),
        ("E Prüfplan", "„Neu“ → Artikel TEST-1, Stand A. Zeile: 1 | Bohrung | Passung | 20 | X7. Zeile: 2 | Länge | (Art leer) | 50.", "Spalte N rot: „Bohrungs-Toleranzfeld X wird nicht unterstützt“ bzw. „Art fehlt“"),
        ("E Prüfplan", "„Speichern“ versuchen.", "Wird abgelehnt, Liste der Fehler; erste Fehlerzeile ist markiert. Danach X7 → H7, Art → Maß, Abmaße ±0,1 → Speichern klappt"),
        ("E Prüfplan", "Blatt ISO-Rechner: 40 und f7.", "Abmaß oben −0,025, unten −0,050, Höchstmaß 39,975, Mindestmaß 39,950"),
        ("F Messung", "„Messung starten / öffnen“ → J1567/5.", "Abgeschlossenes Protokoll öffnet sich nur zum Ansehen (Ergebnis i.O.)"),
        ("F Messung", "„Messung starten / öffnen“ → J1567/40 → Artikel DEMO-200 → Auftrag A-1 → Kunde Test GmbH → Standard-Messmittel MS-01.", "8 Merkmale geladen, Modus BEARBEITEN, Cursor in Istwert Nr 10"),
        ("F Messung", "Istwerte (Spalte H) und Messmittel (Spalte J) eintragen: 10: 80,1 · 20: 39,96 (BM-02) · 30: 30,025 (IM-01) · 40: 12,02 (BM-01) · 50: 60,05 (MS-99) · 60: 0,01 (MU-01) · 70: 1,2 (RA-01) · 80: io (GL-08).", "Bewertung: alle grün, nur Nr 30 rot n.i.O.; Spalte K bei Nr 30 orange"),
        ("F Messung", "„Zwischenspeichern“.", "Fehler: „Messmittel MS-99: Kalibrierung abgelaufen“ – nichts wird gespeichert"),
        ("F Messung", "Nr 50 Messmittel auf MS-02 ändern → „Zwischenspeichern“, dann „Schließen“. Blatt Offene Protokolle → Aktualisieren.", "Statuszeile unten „zwischengespeichert“. Liste: J1567/40 „In Arbeit“, 8 Merkmale, 8 gemessen, 1 n.i.O., 1 offen"),
        ("F Messung", "Zeile J1567/40 anklicken → „Markierte Zeile öffnen“ → „Abschließen“.", "Frage „1 Merkmal(e) n.i.O. ohne Entscheidung – an Freigabe übergeben?“ → Ja → Status „Wartet auf Freigabe“, NUR ANSEHEN"),
        ("G Freigabe", "J1567/40 öffnen. Nr 30: Entscheidung Sonderfreigabe (ohne Begründung) → „Abschließen“.", "Fehler „Begründung für Sonderfreigabe fehlt“"),
        ("G Freigabe", "Begründung eintragen → „Abschließen“ → Ja.", "Ergebnis „i.O. mit Sonderfreigabe“, Meldung mit PDF-Pfad; Entschieden von/am ausgefüllt"),
        ("G Freigabe", "PDF im Ordner PDF\\2026\\DEMO-200 öffnen.", "Kopf mit Produktnummer, Artikelnummer, Auftragsnummer, Kunde; Tabelle quer auf Seitenbreite; Fußzeile Seite x von y"),
        ("G Freigabe", "J1567/40 erneut öffnen.", "Nur ansehen, Hinweis „abgeschlossen“"),
        ("G Freigabe", "Offene Protokolle → J1567/31 öffnen. Nr 2 (rot): Ausschuss + Begründung → Abschließen.", "Ergebnis „Ausschuss“, Protokoll verschwindet aus offenen Protokollen"),
        ("H Nacharbeit", "J1567/32 öffnen.", "Frage „nachgearbeitet – jetzt nachmessen?“ → Ja → Nr 13 Istwert leer, Vorwert 22,13, Bemerkung „Nacharbeit …“"),
        ("H Nacharbeit", "Nr 13 Istwert 22,05 → Abschließen.", "Ergebnis i.O."),
        ("H Nacharbeit", "J1567/33 öffnen (In Arbeit, halb gemessen). Fehlende Werte aus der Mitte der Toleranz eintragen, Nr 11 „nio“ → Entscheidung Nacharbeit → Abschließen.", "Status „Nacharbeit“; danach wie oben nachmessen möglich"),
        ("I Rechte", "Stammdaten\\Stammdaten.xlsx öffnen, bei dir das x bei Freigeber entfernen, speichern und schließen. Im Programm „Anmeldung prüfen“.", "Rollen ohne Freigeber"),
        ("I Rechte", "Neues Protokoll J1567/41 (DEMO-100): Nr 1 = 125 → Sonderfreigabe + Begründung → Zwischenspeichern.", "Fehler „Sonderfreigabe darf nur die Freigabe-Gruppe entscheiden“. Nacharbeit dagegen geht"),
        ("I Rechte", "x bei Freigeber wieder setzen (speichern, schließen), „Anmeldung prüfen“.", "Rolle Freigeber wieder da"),
        ("J Prüfplan-Schutz", "Prüfplan DEMO-200 Stand A laden.", "Hinweis „abgeschlossen (nur lesen)“, weil J1567/40 diesen Stand verwendet"),
        ("K Mehrbenutzer (optional, 2 PCs)", "PC 1 öffnet J1567/41, PC 2 öffnet dasselbe.", "PC 2: „wird gerade von … bearbeitet – nur ansehen?“"),
        ("K Mehrbenutzer (optional, 2 PCs)", "Beide laden Prüfplan TEST-1, ändern etwas, beide speichern.", "Der Zweite bekommt „von jemand anderem geändert“"),
        ("L Abschluss", "Excel schließen (nicht speichern nötig) und Programm neu öffnen.", "Keine Speicherfrage; Startblatt wieder aktuell; offene Sperren freigegeben"),
    ]


def main():
    cpk = cpk_tabelle()
    st = schritte(cpk)
    # Markdown
    md = ["# Großer Test · Messprotokoll 1.0.0", "",
          "Alle Funktionen in einem Durchgang. Dauer etwa 60–90 Minuten.",
          "Abhaken in `Testliste.xlsx` (liegt im ZIP) – Spalte „OK?“ und Notizen.", "",
          "**Vorher:** Windows-PC mit Microsoft 365 Excel. Testordner verwenden, nicht die echten Daten.", ""]
    bereich = None
    for i, (b, s, e) in enumerate(st, start=1):
        if b != bereich:
            md += ["", f"## {b}", ""]
            bereich = b
        md.append(f"{i}. {s}  \n   → **Erwartet:** {e}")
    md += ["", "## Erwartete Cpk-Werte (Testdaten DEMO-100, vor dem Test)", "",
           "| Nr | Merkmal | n | Cpk | Bewertung |", "|---|---|---|---|---|"]
    for nr, name, n, v, t in cpk:
        md.append(f"| {nr} | {name} | {n} | {v:.2f} | {t} |")
    md += ["", "Fehler bitte mit Schrittnummer und Foto der Meldung melden."]
    (BASIS / "TESTPLAN.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    # Excel-Abhakliste
    wb = Workbook()
    ws = wb.active
    ws.title = "Testliste"
    ws["A1"] = "Großer Test · Messprotokoll 1.0.0"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    kopf = [("Nr", 5), ("Bereich", 18), ("Schritt", 70), ("Erwartet", 60), ("OK?", 8), ("Notiz", 40)]
    for i, (t, b) in enumerate(kopf, start=1):
        c = ws.cell(3, i, t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        ws.column_dimensions[c.column_letter].width = b
    for i, (b, s, e) in enumerate(st, start=1):
        for sp, w in enumerate([i, b, s, e], start=1):
            c = ws.cell(3 + i, sp, w)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"ja,nein"', allow_blank=True)
    dv.add(f"E4:E{3 + len(st)}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A4"
    c2 = wb.create_sheet("Cpk erwartet")
    for i, t in enumerate(["Nr", "Merkmal", "n", "Cpk", "Bewertung"], start=1):
        c2.cell(1, i, t).font = Font(bold=True)
    for z, (nr, name, n, v, t) in enumerate(cpk, start=2):
        for sp, w in enumerate([nr, name, n, round(v, 2), t], start=1):
            c2.cell(z, sp, w)
    c2.column_dimensions["B"].width = 26
    wb.save(BASIS / "dateien" / "Testliste.xlsx")
    print(f"Testplan: {len(st)} Schritte, {len(cpk)} Cpk-Werte")


if __name__ == "__main__":
    main()
