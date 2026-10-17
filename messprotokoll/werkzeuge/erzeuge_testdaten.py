"""Testdaten für den großen Test: Prüfpläne, Messmittel, 33 Protokolle.

Ergebnis: dateien/testdaten/ (gleiche Ordnerstruktur wie auf dem Netzlaufwerk)
"""
import random
import shutil
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

import erzeuge_dateien
import protokoll_csv as P

BASIS = Path(__file__).resolve().parent.parent
ZIEL = BASIS / "dateien" / "testdaten"

# Nr, Merkmal, Art, Nennmaß, Passung, Abmaß oben, Abmaß unten, OGW, UGW, Messmittel-Typ, Kritisch
DEMO_100 = [
    ("1", "Gesamtlänge", "Maß", 120, None, 0.1, -0.1, 120.1, 119.9, "Messschieber", "nein"),
    ("2", "Ø Lagersitz", "Passung", 20, "k6", None, None, 20.015, 20.002, "Bügelmessschraube", "ja"),
    ("3", "Ø Wellenabsatz", "Passung", 25, "h6", None, None, 25.0, 24.987, "Bügelmessschraube", "nein"),
    ("4", "Ø Querbohrung", "Passung", 8, "H7", None, None, 8.015, 8.0, "Grenzlehrdorn", "nein"),
    ("5", "Länge Absatz", "Maß", 35, "2768-m", None, None, 35.3, 34.7, "Messschieber", "nein"),
    ("6", "Fase 1x45°", "Maß", 1, "2768-m", None, None, 1.1, 0.9, "Messschieber", "nein"),
    ("7", "Passfedernut Breite", "Passung", 6, "N9", None, None, 6.0, 5.97, "Endmaß", "nein"),
    ("8", "Rundlauf Lagersitz", "Form/Lage", None, None, 0.02, None, 0.02, None, "Messuhr", "ja"),
    ("9", "Ebenheit Stirnfläche", "Form/Lage", None, None, 0.05, None, 0.05, None, "Messuhr", "nein"),
    ("10", "Rauheit Ra Lagersitz", "Oberfläche", None, None, 0.8, None, 0.8, None, "Rauheitsmessgerät", "nein"),
    ("11", "Gewinde M10", "Attribut", None, None, None, None, None, None, "Gewindelehre", "nein"),
    ("12", "Ø Einstich", "Maß", 17, None, 0, -0.1, 17.0, 16.9, "Messschieber", "nein"),
    ("13", "Länge Lagersitz", "Maß", 22, None, 0.1, 0, 22.1, 22.0, "Messschieber", "nein"),
    ("14", "Koaxialität", "Form/Lage", None, None, 0.03, None, 0.03, None, "Messuhr", "nein"),
    ("15", "Radius Übergang", "Maß", 2, "2768-m", None, None, 2.1, 1.9, "Radiuslehre", "nein"),
]
# Messverhalten je Merkmal: (Mittelwert, Streuung) – erzeugt unterschiedliche Cpk-Werte
VERHALTEN = {
    "1": (120.0, 0.02), "2": (20.0085, 0.0015), "3": (24.994, 0.0028), "4": (8.007, 0.0012),
    "5": (35.05, 0.06), "6": (1.0, 0.03), "7": (5.985, 0.004), "8": (0.008, 0.003),
    "9": (0.02, 0.006), "10": (0.5, 0.08), "12": (16.96, 0.012), "13": (22.06, 0.014),
    "14": (0.012, 0.005), "15": (2.0, 0.025),
}
MESSMITTEL_ZU = {"Messschieber": "MS-01", "Bügelmessschraube": "BM-01", "Grenzlehrdorn": "LD-08",
                 "Endmaß": "EM-01", "Messuhr": "MU-01", "Rauheitsmessgerät": "RA-01",
                 "Gewindelehre": "GL-10", "Radiuslehre": "RL-01"}

DEMO_200 = [
    ("10", "Außendurchmesser Flansch", "Maß", 80, "2768-m", None, None, 80.3, 79.7, "Messschieber", "nein"),
    ("20", "Ø Zentrierung", "Passung", 40, "f7", None, None, 39.975, 39.95, "Bügelmessschraube", "ja"),
    ("30", "Ø Bohrung", "Passung", 30, "H7", None, None, 30.021, 30.0, "Innenmessschraube", "ja"),
    ("40", "Dicke", "Maß", 12, None, 0.05, -0.05, 12.05, 11.95, "Bügelmessschraube", "nein"),
    ("50", "Lochkreis", "Maß", 60, None, 0.1, -0.1, 60.1, 59.9, "Messschieber", "nein"),
    ("60", "Planlauf", "Form/Lage", None, None, 0.03, None, 0.03, None, "Messuhr", "nein"),
    ("70", "Rauheit Ra Dichtfläche", "Oberfläche", None, None, 1.6, None, 1.6, None, "Rauheitsmessgerät", "nein"),
    ("80", "Gewinde 4x M8", "Attribut", None, None, None, None, None, None, "Gewindelehre", "nein"),
]

# Nr, Bezeichnung, Art, Messbereich, Auflösung, Standort, Letzte Kalibrierung, Intervall
MESSMITTEL = [
    ("MS-01", "Messschieber digital 150 mm", "Messschieber", "0-150", 0.01, "Dreherei", datetime(2026, 9, 1), 12),
    ("MS-02", "Messschieber 300 mm", "Messschieber", "0-300", 0.01, "Fräserei", datetime(2026, 8, 15), 12),
    ("MS-03", "Messschieber Werkstatt (bald fällig)", "Messschieber", "0-150", 0.01, "Dreherei", datetime(2025, 11, 5), 12),
    ("MS-99", "Messschieber alt (abgelaufen)", "Messschieber", "0-150", 0.02, "Lager", datetime(2024, 3, 1), 12),
    ("BM-01", "Bügelmessschraube 0-25", "Bügelmessschraube", "0-25", 0.001, "Messraum", datetime(2026, 7, 1), 12),
    ("BM-02", "Bügelmessschraube 25-50", "Bügelmessschraube", "25-50", 0.001, "Messraum", datetime(2026, 7, 1), 12),
    ("IM-01", "Innenmessschraube 25-50", "Innenmessschraube", "25-50", 0.001, "Messraum", datetime(2026, 6, 1), 12),
    ("LD-08", "Grenzlehrdorn Ø8 H7", "Lehre", "8", None, "Messraum", datetime(2026, 2, 1), 24),
    ("EM-01", "Endmaßsatz", "Endmaß", "0,5-100", None, "Messraum", datetime(2026, 1, 15), 24),
    ("MU-01", "Messuhr 0,001", "Messuhr", "0-1", 0.001, "Messraum", datetime(2026, 5, 1), 12),
    ("RA-01", "Rauheitsmessgerät", "Rauheit", "Ra 0,05-40", 0.01, "Messraum", datetime(2026, 4, 1), 12),
    ("GL-10", "Gewindelehrring M10", "Lehre", "M10", None, "Dreherei", datetime(2026, 3, 1), 24),
    ("GL-08", "Gewindelehrdorn M8", "Lehre", "M8", None, "Fräserei", datetime(2026, 3, 1), 24),
    ("RL-01", "Radiuslehrensatz", "Lehre", "1-7", None, "Dreherei", datetime(2026, 2, 1), 24),
    ("RL-02", "Radiuslehrensatz alt (ohne Kalibrierdaten)", "Lehre", "1-7", None, "Lager", None, None),
]


def pruefplan(artikel, bezeichnung, merkmale, datei):
    wb = Workbook()
    k = wb.active
    k.title = "Kopf"
    for z, (a, b) in enumerate([
        ("Artikelnummer", artikel), ("Bezeichnung", bezeichnung), ("Zeichnungsstand", "A"),
        ("Basis-Stand", None), ("Erstellt von", "Testdaten"), ("Erstellt am", datetime(2026, 7, 1, 9, 0)),
        ("Geändert von", "Testdaten"), ("Geändert am", datetime(2026, 7, 1, 9, 0)),
        ("Entfallene Merkmale", None), ("Anzahl Merkmale", len(merkmale)), ("Programmversion", "1.0.0"),
    ], start=1):
        k.cell(z, 1, a)
        k.cell(z, 2, b)
    m = wb.create_sheet("Merkmale")
    kopf = ["Nr", "Merkmal", "Art", "Nennmaß", "Passung / Toleranz", "Abmaß oben", "Abmaß unten",
            "OGW", "UGW", "Messmittel", "Kritisch", "Änderung", "Bemerkung"]
    for i, t in enumerate(kopf, start=1):
        m.cell(1, i, t)
    for z, e in enumerate(merkmale, start=2):
        for i, w in enumerate(e[:9], start=1):
            m.cell(z, i, w)
        m.cell(z, 10, e[9])
        m.cell(z, 11, e[10])
    datei.parent.mkdir(parents=True, exist_ok=True)
    wb.save(datei)
    (datei.parent / "aktueller_Stand.txt").write_bytes(b"A\r\n")


def bewerte(art, ist, ogw, ugw):
    if art == "Attribut":
        return {"i.O.": "i.O.", "n.i.O.": "n.i.O."}.get(ist, "")
    if ist is None:
        return ""
    if ogw is not None and ist > ogw + 1e-7:
        return "n.i.O."
    if ugw is not None and ist < ugw - 1e-7:
        return "n.i.O."
    return "i.O."


def protokoll(nr, status, start, rnd, nio_plan=None):
    """Erzeugt Kopf und Merkmale eines DEMO-100-Protokolls."""
    prod = f"J1567/{nr}"
    merkmale = []
    for e in DEMO_100:
        mnr, name, art, nenn, pas, _o, _u, ogw, ugw, typ, krit = e
        if art == "Attribut":
            ist = "i.O."
        else:
            mittel, streu = VERHALTEN[mnr]
            # nur Werte innerhalb der Toleranz (n.i.O. kommen gezielt über nio_plan)
            for _ in range(1000):
                ist = round(rnd.gauss(mittel, streu), 4 if streu < 0.01 else 3)
                if art in ("Form/Lage", "Oberfläche"):
                    ist = abs(ist)
                if bewerte(art, ist, ogw, ugw) == "i.O.":
                    break
        merkmale.append({"Nr": mnr, "Merkmal": name, "Art": art, "Nennmaß": nenn, "Passung / Toleranz": pas,
                         "OGW": ogw, "UGW": ugw, "Istwert": ist, "Messmittel": MESSMITTEL_ZU[typ],
                         "Kritisch": krit, "Messmittel-Typ": typ})
    if nio_plan:
        for mnr, wert, ent, grund in nio_plan:
            m = next(x for x in merkmale if x["Nr"] == mnr)
            m["Istwert"] = wert
            if ent:
                m["Entscheidung"] = ent
                m["Begründung"] = grund
                m["Entschieden von"] = "Demo Freigabe" if ent != "Nacharbeit" else "Demo Prüfer"
                m["Entschieden am"] = start + timedelta(minutes=50)
    ergebnis = None
    for m in merkmale:
        b = bewerte(m["Art"], m["Istwert"], m["OGW"], m["UGW"])
        m["Bewertung"] = b or None
        if b != "n.i.O.":
            for f in ("Entscheidung", "Begründung", "Entschieden von", "Entschieden am"):
                m.pop(f, None)
    if status == "Abgeschlossen":
        ents = [m.get("Entscheidung") for m in merkmale if m["Bewertung"] == "n.i.O."]
        ergebnis = "Ausschuss" if "Ausschuss" in ents else ("i.O. mit Sonderfreigabe" if ents else "i.O.")
    kopf = {"Produktnummer": prod, "Artikelnummer": "DEMO-100", "Bezeichnung": "Antriebswelle",
            "Auftragsnummer": f"A-2026-{400 + nr // 10}", "Kunde": "Muster Maschinenbau GmbH",
            "Prüfplan-Stand": "A", "Status": status, "Ergebnis": ergebnis, "Prüfer": "Demo Prüfer",
            "Begonnen": start, "Abgeschlossen": start + timedelta(hours=1) if status == "Abgeschlossen" else None,
            "Geändert von": "Demo Prüfer", "Geändert am": start + timedelta(hours=1), "Programmversion": "1.0.0"}
    return prod, kopf, merkmale


def main():
    if ZIEL.exists():
        shutil.rmtree(ZIEL)
    rnd = random.Random(1567)
    (ZIEL / "Stammdaten").mkdir(parents=True)
    erzeuge_dateien.stammdaten(ZIEL / "Stammdaten" / "Stammdaten.xlsx", MESSMITTEL)
    pruefplan("DEMO-100", "Antriebswelle", DEMO_100, ZIEL / "Pruefplaene" / "DEMO-100" / "DEMO-100_Stand-A.xlsx")
    pruefplan("DEMO-200", "Flansch", DEMO_200, ZIEL / "Pruefplaene" / "DEMO-200" / "DEMO-200_Stand-A.xlsx")

    besonders = {
        7: [("2", 20.0162, "Sonderfreigabe", "Kunde informiert, Lager passt laut Montageversuch")],
        19: [("5", 35.42, "Sonderfreigabe", "Funktion nicht beeinträchtigt (Abstimmung Konstruktion)")],
        23: [("3", 24.982, "Ausschuss", "Untermaß, Nacharbeit nicht möglich")],
    }
    start = datetime(2026, 8, 3, 7, 30)
    for nr in range(1, 31):
        prod, kopf, merkmale = protokoll(nr, "Abgeschlossen", start, rnd, besonders.get(nr))
        rel = f"Daten\\Archiv\\2026\\DEMO-100\\{prod.replace('/', '_')}.csv"
        P.schreiben(ZIEL / rel.replace("\\", "/"), kopf, merkmale)
        (ZIEL / "System" / "Index").mkdir(parents=True, exist_ok=True)
        (ZIEL / "System" / "Index" / f"{prod.replace('/', '_')}.txt").write_bytes((rel + "\r\n").encode("cp1252"))
        start += timedelta(days=2, hours=rnd.randint(0, 6))

    offen = [
        (31, "Wartet auf Freigabe", "Daten\\Freigabe", [("2", 20.0171, None, None)]),
        (32, "Nacharbeit", "Daten\\Entwurf", [("13", 22.13, "Nacharbeit", "Länge zu lang – nachdrehen")]),
        (33, "In Arbeit", "Daten\\Entwurf", None),
    ]
    for nr, status, ordner, plan in offen:
        prod, kopf, merkmale = protokoll(nr, status, datetime(2026, 10, 14, 8, 0), rnd, plan)
        if nr == 33:
            for m in merkmale[6:]:
                m["Istwert"] = None
                m["Bewertung"] = None
        rel = f"{ordner}\\{prod.replace('/', '_')}.csv"
        P.schreiben(ZIEL / rel.replace("\\", "/"), kopf, merkmale)
        (ZIEL / "System" / "Index" / f"{prod.replace('/', '_')}.txt").write_bytes((rel + "\r\n").encode("cp1252"))
    print("Testdaten erzeugt:", sum(1 for _ in ZIEL.rglob("*") if _.is_file()), "Dateien")


if __name__ == "__main__":
    main()
